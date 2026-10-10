"""Loads the Switzerland database (Excel) into plain Python structures."""
import os
from openpyxl import load_workbook

DB_PATH = os.environ.get("DB_PATH", os.path.join(os.path.dirname(__file__), "Switzerland_Database.xlsx"))

ZRH = "Zurich Airport (ZRH)"
GVA = "Geneva Airport (GVA)"
ZERMATT_KEY = "Zermatt (park at Tasch)"


def loc_key(name: str) -> str:
    """Map a stay-city name to the location name used in the Drive Distances sheet."""
    return ZERMATT_KEY if name == "Zermatt" else name


def load_db(path: str = DB_PATH) -> dict:
    wb = load_workbook(path, data_only=True)

    cities = {}
    for row in wb["Stay Cities"].iter_rows(min_row=2, values_only=True):
        if row[0]:
            cities[row[0]] = {"why": row[1], "best_for": row[2], "airport": row[3]}

    ws = wb["Sightseeing"]
    sights = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row[0]:
            continue
        sights.append({
            "place": row[0], "city": row[1], "type": row[2], "time": row[3],
            "adult_inr_sheet": row[4], "child_inr_sheet": row[5],
            "adult_chf": float(row[6] or 0), "child_chf": float(row[7] or 0),
            "status": row[8], "note": row[9],
        })

    # Assumption block: columns L (label) and M (value), rows 2-5
    assume = {}
    for r in range(2, 6):
        assume[ws.cell(r, 12).value] = ws.cell(r, 13).value
    cfg = {
        "sheet_fx": float(assume["CHF to INR rate (reference only)"]),
        "vignette_chf": float(assume["Swiss motorway vignette (CHF)"]),
        "parking_chf_per_night": float(assume["Parking per night (CHF)"]),
        "service_charge_inr": float(assume["Romford service charge per adult (INR)"]),
    }

    dist = {}
    for a, b, km, _txt, mins, _note in wb["Drive Distances"].iter_rows(min_row=2, values_only=True):
        if a and b:
            dist[(a, b)] = (km, int(mins))
            dist[(b, a)] = (km, int(mins))
    return {"cities": cities, "sights": sights, "dist": dist, "cfg": cfg}


def drive(db: dict, a: str, b: str):
    """Return (km, minutes) between two location names, or None when same place."""
    if a == b:
        return None
    return db["dist"][(a, b)]


def fmt_minutes(m: int) -> str:
    h, mm = divmod(int(m), 60)
    return f"{h} h {mm:02d} min" if h else f"{mm} min"
